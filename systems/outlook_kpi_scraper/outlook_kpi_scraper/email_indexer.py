"""
Email RAG Indexer — System 2 of the Hybrid Architecture.

Reads emails from Outlook via COM and indexes them into a local ChromaDB
vector store using OpenAI text-embedding-3-small.  Each email becomes one
or more documents (body chunk + per-attachment text chunk).

Usage (standalone):
    python -m outlook_kpi_scraper.email_indexer \
        --mailbox "Chip Ridge" --folders "Inbox,Sent Items,Junk Email" \
        --days 30 --max 5000

After indexing, the query_agent module can search the collection.
"""

import argparse
import hashlib
import logging
import os
import re
import time
from datetime import datetime

import chromadb
from chromadb.config import Settings as ChromaSettings

# OutlookReader uses win32com (Windows-only).  Import lazily so that Linux
# deployments (Railway) can still use _get_collection / query functions
# without crashing at import time.
try:
    from outlook_kpi_scraper.outlook_reader import OutlookReader
except ImportError:
    OutlookReader = None  # type: ignore[misc,assignment]

from outlook_kpi_scraper.utils import load_env

log = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

COLLECTION_NAME = "chip_emails"
CHROMA_DIR = os.environ.get(
    "CHROMA_PERSIST_DIR",
    os.path.join(os.path.dirname(__file__), "..", "data", "chromadb"),
)
EMBED_MODEL = "text-embedding-3-small"
MAX_CHUNK_CHARS = 6000  # ~1 500 tokens — fits embedding model context
OVERLAP_CHARS = 400


# ---------------------------------------------------------------------------
# Embedding function (OpenAI)
# ---------------------------------------------------------------------------

class OpenAIEmbedder(chromadb.EmbeddingFunction):
    """Wraps OpenAI text-embedding-3-small for ChromaDB."""

    def __init__(self, api_key: str, model: str = EMBED_MODEL):
        from openai import OpenAI
        self._client = OpenAI(api_key=api_key)
        self._model = model

    def __call__(self, input: list[str]) -> list[list[float]]:
        # Batch embed (max 2048 inputs per call for OpenAI)
        results = []
        batch_size = 512
        for i in range(0, len(input), batch_size):
            batch = input[i : i + batch_size]
            resp = self._client.embeddings.create(model=self._model, input=batch)
            results.extend([d.embedding for d in resp.data])
        return results


# ---------------------------------------------------------------------------
# Chunking helpers
# ---------------------------------------------------------------------------

def _chunk_text(text: str, max_chars: int = MAX_CHUNK_CHARS,
                overlap: int = OVERLAP_CHARS) -> list[str]:
    """Split *text* into overlapping chunks respecting paragraph boundaries."""
    if len(text) <= max_chars:
        return [text]
    chunks = []
    start = 0
    while start < len(text):
        end = start + max_chars
        # Try to break at a paragraph boundary
        if end < len(text):
            newline_pos = text.rfind("\n\n", start + max_chars // 2, end)
            if newline_pos != -1:
                end = newline_pos
        chunks.append(text[start:end].strip())
        start = end - overlap
    return [c for c in chunks if c]


def _stable_id(entry_id: str, chunk_idx: int, source: str = "body",
               *, fallback_key: str = "") -> str:
    """Generate a deterministic doc ID so re-indexing is idempotent.

    If *entry_id* is empty, *fallback_key* (e.g. subject+sender+date)
    is used to avoid collisions.
    """
    key = entry_id or fallback_key or "unknown"
    raw = f"{key}::{source}::{chunk_idx}"
    return hashlib.sha256(raw.encode()).hexdigest()[:24]


def _strip_signature(body: str) -> str:
    """Remove common email signatures and quoted replies."""
    # Strip quoted replies
    for marker in [
        r"\n-{3,}\s*Original Message",
        r"\nOn .{10,80} wrote:",
        r"\nFrom: .{5,80}\nSent:",
        r"\n_{3,}",
    ]:
        m = re.search(marker, body, re.IGNORECASE)
        if m:
            body = body[: m.start()]
    # Strip trailing signature
    lines = body.rstrip().split("\n")
    # If last 5 lines are very short (signature), drop them
    if len(lines) > 8:
        tail = lines[-5:]
        if all(len(l.strip()) < 80 for l in tail):
            # Check if there's a "--" or similar separator
            for i in range(len(lines) - 6, len(lines)):
                if re.match(r"^[-–—]{2,}$", lines[i].strip()):
                    body = "\n".join(lines[:i])
                    break
    return body.strip()


# ---------------------------------------------------------------------------
# Core indexing
# ---------------------------------------------------------------------------

def _get_collection(api_key: str):
    """Return (client, collection) — creates them if needed."""
    os.makedirs(CHROMA_DIR, exist_ok=True)
    client = chromadb.PersistentClient(
        path=CHROMA_DIR,
        settings=ChromaSettings(anonymized_telemetry=False),
    )
    embedder = OpenAIEmbedder(api_key)
    collection = client.get_or_create_collection(
        name=COLLECTION_NAME,
        embedding_function=embedder,
        metadata={"hnsw:space": "cosine"},
    )
    return client, collection


def index_messages(messages: list[dict], api_key: str, *,
                   skip_existing: bool = True) -> dict:
    """Index a list of message dicts into ChromaDB.

    Returns summary dict with counts.
    """
    _client, collection = _get_collection(api_key)

    stats = {
        "total_messages": len(messages),
        "indexed_docs": 0,
        "skipped_existing": 0,
        "skipped_empty": 0,
        "embed_cost_estimate": 0.0,
    }

    docs_batch: list[str] = []
    ids_batch: list[str] = []
    metas_batch: list[dict] = []

    for msg in messages:
        entry_id = msg.get("entry_id", "")
        subject = msg.get("subject", "") or "(no subject)"
        sender = msg.get("sender_email", "") or msg.get("sender_name", "")
        received = msg.get("received_dt", "")[:10]
        folder = msg.get("source_folder", "Inbox")
        recipients = msg.get("recipients_to", "")
        body = _strip_signature(msg.get("body", "") or "")
        message_id = msg.get("internet_message_id", "") or ""
        thread_topic = msg.get("conversation_topic", "") or ""

        # Extract sender domain for filtering
        sender_domain = ""
        if "@" in sender:
            sender_domain = sender.rsplit("@", 1)[-1].lower()

        if not body or len(body.strip()) < 30:
            stats["skipped_empty"] += 1
            continue

        # Build text: subject + metadata header + body
        header = (
            f"Subject: {subject}\n"
            f"From: {sender}\n"
            f"Date: {received}\n"
            f"Folder: {folder}\n"
        )
        if recipients:
            header += f"To: {recipients}\n"
        header += f"Attachments: {msg.get('attachment_names', '')}\n"

        full_text = header + "\n" + body

        # Build a fallback key for messages without an entry_id
        fallback_key = f"{subject}|{sender}|{received}|{message_id}"

        chunks = _chunk_text(full_text)
        for ci, chunk in enumerate(chunks):
            doc_id = _stable_id(entry_id, ci, fallback_key=fallback_key)

            if skip_existing:
                try:
                    existing = collection.get(ids=[doc_id])
                    if existing and existing["ids"]:
                        stats["skipped_existing"] += 1
                        continue
                except Exception:
                    pass

            meta = {
                "entry_id": entry_id,
                "subject": subject[:200],
                "sender": sender[:120],
                "sender_domain": sender_domain,
                "date": received,
                "folder": folder,
                "message_id": message_id[:200],
                "thread_topic": thread_topic[:200],
                "has_attachments": str(msg.get("has_attachments", False)),
                "attachment_names": (msg.get("attachment_names", "") or "")[:300],
                "chunk_index": ci,
                "total_chunks": len(chunks),
            }

            docs_batch.append(chunk)
            ids_batch.append(doc_id)
            metas_batch.append(meta)

    # Deduplicate within the batch (ChromaDB rejects duplicate IDs in one upsert)
    seen_ids: set[str] = set()
    deduped_docs, deduped_ids, deduped_metas = [], [], []
    for doc, did, meta in zip(docs_batch, ids_batch, metas_batch):
        if did not in seen_ids:
            seen_ids.add(did)
            deduped_docs.append(doc)
            deduped_ids.append(did)
            deduped_metas.append(meta)
        else:
            stats["skipped_existing"] += 1

    # Batch upsert into ChromaDB (chromadb handles embedding internally)
    BATCH = 200
    total_chars = 0
    for i in range(0, len(deduped_docs), BATCH):
        batch_docs = deduped_docs[i : i + BATCH]
        batch_ids = deduped_ids[i : i + BATCH]
        batch_metas = deduped_metas[i : i + BATCH]
        collection.upsert(
            documents=batch_docs,
            ids=batch_ids,
            metadatas=batch_metas,
        )
        total_chars += sum(len(d) for d in batch_docs)
        stats["indexed_docs"] += len(batch_docs)

    # Rough cost estimate: ~$0.02 / 1M tokens, ~4 chars per token
    tokens_est = total_chars / 4
    stats["embed_cost_estimate"] = round(tokens_est / 1_000_000 * 0.02, 4)

    log.info(
        "Indexing complete: indexed=%d, skipped_existing=%d, skipped_empty=%d, "
        "cost≈$%.4f",
        stats["indexed_docs"],
        stats["skipped_existing"],
        stats["skipped_empty"],
        stats["embed_cost_estimate"],
    )
    return stats


def get_collection_stats(api_key: str) -> dict:
    """Return basic stats about the current ChromaDB collection."""
    _client, collection = _get_collection(api_key)
    count = collection.count()
    chroma_abs = os.path.abspath(CHROMA_DIR)
    return {
        "collection": COLLECTION_NAME,
        "document_count": count,
        "chromadb_path": chroma_abs,
    }


# ---------------------------------------------------------------------------
# Attachment text extraction (conservative — PDF / XLSX / CSV / DOCX only)
# ---------------------------------------------------------------------------

_SAFE_TEXT_EXTS = {".pdf", ".xlsx", ".xls", ".csv", ".docx", ".txt"}
_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff"}


def _extract_attachment_text(reader: "OutlookReader", msg: dict,
                              tmp_dir: str) -> list[dict]:
    """Extract text from safe attachment types for a single message.

    Returns a list of {filename, text} dicts.  Skips images entirely.
    """
    entry_id = msg.get("entry_id", "")
    if not entry_id or not msg.get("has_attachments"):
        return []

    raw_item = reader.get_raw_item(entry_id)
    if raw_item is None:
        return []

    att_count = 0
    try:
        att_count = raw_item.Attachments.Count
    except Exception:
        return []
    if att_count == 0:
        return []

    results = []
    msg_dir = os.path.join(tmp_dir, hashlib.md5(entry_id.encode()).hexdigest()[:12])
    os.makedirs(msg_dir, exist_ok=True)

    for idx in range(1, att_count + 1):
        try:
            att = raw_item.Attachments.Item(idx)
            fname = getattr(att, "FileName", f"att_{idx}") or f"att_{idx}"
            ext = os.path.splitext(fname)[1].lower()
            if ext in _IMAGE_EXTS or ext not in _SAFE_TEXT_EXTS:
                continue
            size = getattr(att, "Size", 0)
            if size > 8 * 1024 * 1024:  # skip > 8 MB
                continue

            dest = os.path.join(msg_dir, fname)
            att.SaveAsFile(dest)
            if not os.path.exists(dest):
                continue

            text = _read_file_text(dest, ext)
            if text and len(text.strip()) >= 50:
                results.append({"filename": fname, "text": text[:MAX_CHUNK_CHARS * 3]})
        except Exception as exc:
            log.debug("Attachment extraction failed idx=%d: %s", idx, exc)
    return results


def _read_file_text(path: str, ext: str) -> str:
    """Read plain text from a file based on extension."""
    try:
        if ext == ".txt":
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        elif ext == ".csv":
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        elif ext in (".xlsx", ".xls"):
            return _read_excel_text(path, ext)
        elif ext == ".pdf":
            return _read_pdf_text(path)
        elif ext == ".docx":
            return _read_docx_text(path)
    except Exception as exc:
        log.debug("Failed to read %s: %s", path, exc)
    return ""


def _read_excel_text(path: str, ext: str) -> str:
    """Read all cells from an Excel file as text."""
    try:
        import openpyxl
        if ext == ".xlsx":
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets[:3]:  # max 3 sheets
                for row in ws.iter_rows(max_row=200, values_only=True):
                    vals = [str(c) if c is not None else "" for c in row]
                    line = " | ".join(v for v in vals if v)
                    if line.strip():
                        lines.append(line)
            wb.close()
            return "\n".join(lines)
    except ImportError:
        pass
    try:
        import xlrd
        wb = xlrd.open_workbook(path)
        lines = []
        for ws in wb.sheets()[:3]:
            for rx in range(min(ws.nrows, 200)):
                vals = [str(ws.cell_value(rx, cx)) for cx in range(ws.ncols)]
                line = " | ".join(v for v in vals if v)
                if line.strip():
                    lines.append(line)
        return "\n".join(lines)
    except ImportError:
        pass
    return ""


def _read_pdf_text(path: str) -> str:
    """Extract text from a PDF (pypdf first, pdfminer fallback)."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(path)
        pages = []
        for page in reader.pages[:20]:  # max 20 pages
            pages.append(page.extract_text() or "")
        return "\n".join(pages)
    except Exception:
        pass
    try:
        from pdfminer.high_level import extract_text
        return extract_text(path, maxpages=20)
    except Exception:
        pass
    return ""


def _read_docx_text(path: str) -> str:
    """Extract text from a DOCX file."""
    try:
        from docx import Document
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception:
        pass
    return ""


def index_attachments(messages: list[dict], reader: "OutlookReader",
                       api_key: str, *, skip_existing: bool = True) -> dict:
    """Index attachment text from messages into ChromaDB.

    Returns stats dict.
    """
    import tempfile

    _client, collection = _get_collection(api_key)
    tmp_dir = os.path.join(os.path.dirname(__file__), "..", "data", "att_tmp")
    os.makedirs(tmp_dir, exist_ok=True)

    stats = {"att_docs_indexed": 0, "att_skipped": 0, "att_errors": 0}

    docs_batch, ids_batch, metas_batch = [], [], []

    for msg in messages:
        if not msg.get("has_attachments"):
            continue
        try:
            att_texts = _extract_attachment_text(reader, msg, tmp_dir)
        except Exception as exc:
            log.debug("Attachment extraction error: %s", exc)
            stats["att_errors"] += 1
            continue

        for att in att_texts:
            entry_id = msg.get("entry_id", "")
            chunks = _chunk_text(att["text"])
            for ci, chunk in enumerate(chunks):
                doc_id = _stable_id(entry_id, ci, source=f"att:{att['filename']}")
                if skip_existing:
                    try:
                        existing = collection.get(ids=[doc_id])
                        if existing and existing["ids"]:
                            stats["att_skipped"] += 1
                            continue
                    except Exception:
                        pass

                sender = msg.get("sender_email", "") or msg.get("sender_name", "")
                sender_domain = ""
                if "@" in sender:
                    sender_domain = sender.rsplit("@", 1)[-1].lower()

                meta = {
                    "entry_id": entry_id,
                    "subject": (msg.get("subject", "") or "")[:200],
                    "sender": sender[:120],
                    "sender_domain": sender_domain,
                    "date": (msg.get("received_dt", "") or "")[:10],
                    "folder": msg.get("source_folder", ""),
                    "message_id": (msg.get("internet_message_id", "") or "")[:200],
                    "thread_topic": (msg.get("conversation_topic", "") or "")[:200],
                    "has_attachments": "True",
                    "attachment_names": att["filename"],
                    "source_type": "attachment",
                    "chunk_index": ci,
                    "total_chunks": len(chunks),
                }
                # Prefix doc with attachment context
                header = (
                    f"[Attachment: {att['filename']}]\n"
                    f"From email: {msg.get('subject', '')} | {sender} | "
                    f"{(msg.get('received_dt', '') or '')[:10]}\n\n"
                )
                docs_batch.append(header + chunk)
                ids_batch.append(doc_id)
                metas_batch.append(meta)

    # Batch upsert
    BATCH = 200
    for i in range(0, len(docs_batch), BATCH):
        collection.upsert(
            documents=docs_batch[i : i + BATCH],
            ids=ids_batch[i : i + BATCH],
            metadatas=metas_batch[i : i + BATCH],
        )
        stats["att_docs_indexed"] += len(docs_batch[i : i + BATCH])

    # Clean up temp dir
    import shutil
    try:
        shutil.rmtree(tmp_dir, ignore_errors=True)
    except Exception:
        pass

    log.info("Attachment indexing: indexed=%d skipped=%d errors=%d",
             stats["att_docs_indexed"], stats["att_skipped"], stats["att_errors"])
    return stats


# ---------------------------------------------------------------------------
# CLI entry point
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Index Outlook emails into ChromaDB for RAG queries."
    )
    parser.add_argument("--mailbox", required=True, help='Mailbox display name (e.g. "Chip Ridge")')
    parser.add_argument("--folders", default="Inbox",
                        help='Comma-separated folders (e.g. "Inbox,Sent Items,Junk Email")')
    parser.add_argument("--days", type=int, default=30, help="Days to look back")
    parser.add_argument("--max", type=int, default=5000, help="Max messages per folder")
    parser.add_argument("--reindex", action="store_true",
                        help="Force re-index even if doc already exists")
    parser.add_argument("--include-attachments", action="store_true",
                        help="Also extract and index text from PDF/XLSX/CSV/DOCX attachments (slower)")
    args = parser.parse_args()

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(name)s: %(message)s",
    )

    env = load_env()
    api_key = env.get("OPENAI_API_KEY") or os.environ.get("OPENAI_API_KEY")
    if not api_key:
        print("ERROR: OPENAI_API_KEY not set")
        return

    folder_list = [f.strip() for f in args.folders.split(",") if f.strip()]

    print(f"\n{'='*60}")
    print(f"  EMAIL RAG INDEXER")
    print(f"  mailbox={args.mailbox}  folders={folder_list}")
    print(f"  days={args.days}  max={args.max}")
    print(f"  include_attachments={args.include_attachments}")
    print(f"{'='*60}\n")

    t0 = time.time()

    # Fetch from Outlook
    reader = OutlookReader(
        mailbox=args.mailbox, folder=folder_list,
        days=args.days, max_items=args.max,
    )
    messages = reader.fetch_messages()
    print(f"  Fetched {len(messages)} messages from Outlook")

    # Index email bodies
    stats = index_messages(messages, api_key, skip_existing=not args.reindex)

    # Optionally index attachment text
    att_stats = {}
    if args.include_attachments:
        print("  Extracting and indexing attachment text...")
        att_stats = index_attachments(messages, reader, api_key,
                                       skip_existing=not args.reindex)

    duration = time.time() - t0
    print(f"\n{'='*60}")
    print(f"  INDEXING COMPLETE")
    print(f"  bodies: indexed={stats['indexed_docs']}  skipped_existing={stats['skipped_existing']}  "
          f"skipped_empty={stats['skipped_empty']}")
    if att_stats:
        print(f"  attachments: indexed={att_stats.get('att_docs_indexed', 0)}  "
              f"skipped={att_stats.get('att_skipped', 0)}  "
              f"errors={att_stats.get('att_errors', 0)}")
    print(f"  embed_cost≈${stats['embed_cost_estimate']:.4f}")
    print(f"  duration={duration:.1f}s")

    # Show collection stats
    coll_stats = get_collection_stats(api_key)
    print(f"  collection_total={coll_stats['document_count']} docs")
    print(f"  chromadb_path={coll_stats['chromadb_path']}")
    print(f"{'='*60}")


if __name__ == "__main__":
    main()
